"""CSV or Excel bytes -> typed Arrow table with the register's field names and nothing else.

Rules that live here: allow-list ingest, blank cells become null, a field's declared unknown
markers become null, suppression tokens become null plus a typed flag, a wide table of dated
columns is unpivoted to one row per cell when the register says so, and a value that cannot take
its declared type stops the build instead of being coerced.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from typing import TYPE_CHECKING, Any, cast

import openpyxl
import pyarrow as pa
import pyarrow.compute as pc
import pyarrow.csv as pcsv
import xlrd

from .register import (
    CELL_OF_RE,
    CELL_SOURCE,
    COLUMN_HEADER_RE,
    HEADER_SOURCE,
    LAT_SOURCE,
    LON_SOURCE,
    ROW_HEADER_RE,
    SHEET_SOURCE,
    Dataset,
    Field,
)
from .spine import is_spine, read_points, read_shapes

if TYPE_CHECKING:
    from .register import Wide
    from .store import Manifest

# pyarrow.compute takes either and hands back either; pyarrow-stubs 20 often names the wrong one.
type Arr = pa.Array[Any] | pa.ChunkedArray[Any]

ARROW_TYPES = {
    "string": pa.string(),
    "integer": pa.int64(),
    "number": pa.float64(),
    "boolean": pa.bool_(),
    "date": pa.date32(),
    "datetime": pa.timestamp("s"),
}


class NormaliseError(ValueError):
    pass


@dataclass
class Table:
    dataset: Dataset
    manifest: Manifest
    table: pa.Table
    unknown_columns: list[str] = field(default_factory=list)
    suppressed_cells: int = 0
    # Rows a spreadsheet export left short of the header, padded with blank cells to read them.
    short_rows: int = 0
    # Upstream columns the register leaves out on purpose, seen in this file.
    omitted_columns: list[str] = field(default_factory=list)
    # A polygon or line layer's geometry, one WKB value per row in GDA2020, beside the fields.
    geometry: pa.Array[Any] | None = None
    # The place spine layers a point dataset was joined to, with the version of each.
    places: list[dict[str, str]] = field(default_factory=list)
    # (table, (sort, key), permutation) once the build has sorted this table, so it sorts once.
    order: tuple[pa.Table, tuple[tuple[str, ...], tuple[str, ...]], pa.Array[Any] | None] | None = (
        field(default=None, repr=False, compare=False)
    )

    @property
    def rows(self) -> int:
        return self.table.num_rows


def detect_encoding(data: bytes, preferred: str = "") -> str:
    for enc in [preferred, "utf-8-sig", "cp1252"]:
        if not enc:
            continue
        try:
            data.decode(enc)
        except UnicodeDecodeError:
            continue
        return enc
    return "latin-1"


def _header(text: str, delimiter: str = ",") -> list[str]:
    # As written, spaces included: Arrow names the columns this way, so every one is typed as text.
    return next(csv.reader(io.StringIO(text.splitlines()[0]), delimiter=delimiter))


def _delimiter(declared: str) -> str:
    return {"": ",", "tab": "\t"}.get(declared, declared)


def read_csv(
    data: bytes,
    encoding: str,
    delimiter: str = "",
    header_row: int = 1,
    short: list[int] | None = None,
) -> pa.Table:
    """The file as text columns. `short`, when given, receives the count of rows padded."""
    text = data.decode(encoding)
    if header_row > 1:
        text = "".join(text.splitlines(keepends=True)[header_row - 1 :])
    sep = _delimiter(delimiter)
    header = _header(text, sep)
    try:
        t = _arrow_csv(text, header, sep)
    except pa.ArrowInvalid as e:
        # A spreadsheet export leaves the empty cells off the end of a row; pad them back.
        if "Expected" not in str(e) or "columns" not in str(e):
            raise
        padded, n = _pad_rows(text, sep, len(header))
        if short is not None:
            short.append(n)
        t = _arrow_csv(padded, header, sep)
    return _drop_blank_rows(t)


def _arrow_csv(text: str, header: list[str], sep: str) -> pa.Table:
    return pcsv.read_csv(
        io.BytesIO(text.encode("utf-8")),
        read_options=pcsv.ReadOptions(encoding="utf8", block_size=8 << 20),
        parse_options=pcsv.ParseOptions(newlines_in_values=True, delimiter=sep),
        convert_options=pcsv.ConvertOptions(
            column_types={h: pa.string() for h in header},
            strings_can_be_null=False,
            null_values=[],
            quoted_strings_can_be_null=False,
        ),
    )


def _pad_rows(text: str, sep: str, width: int) -> tuple[str, int]:
    out = io.StringIO()
    w = csv.writer(out, delimiter=sep, lineterminator="\n")
    n = 0
    for row in csv.reader(io.StringIO(text), delimiter=sep):
        if len(row) < width:
            n += 1
            w.writerow(row + [""] * (width - len(row)))
        else:
            w.writerow(row)
    return out.getvalue(), n


def _drop_blank_rows(t: pa.Table) -> pa.Table:
    """A spreadsheet exported as CSV ends in rows of separators alone; they are not data."""
    if not t.num_rows or not t.num_columns:
        return t
    blank: Arr | None = None
    for c in t.columns:
        b = pc.equal(pc.utf8_trim_whitespace(pc.fill_null(c, "")), "")  # type: ignore[call-overload, type-var]  # pyarrow-stubs 20 types fill_null as coalesce and takes no Python scalar
        blank = b if blank is None else pc.and_(blank, b)
    return t.filter(pc.invert(blank)) if pc.any(blank).as_py() else t  # type: ignore[arg-type, type-var]  # a table with a column always sets blank


def _cell(v: object) -> str:
    """An Excel cell as the text the publisher would have typed, so typing follows one path."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    return str(v)


def xls_to_xlsx(data: bytes) -> bytes:
    """A legacy Excel 97-2003 workbook as an xlsx one, sheet for sheet and cell for cell.

    Every workbook reader then reads it the same way. Date cells stay dates.
    """
    book = xlrd.open_workbook(file_contents=data)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]  # a new Workbook always has its first sheet
    for sh in book.sheets():
        ws = wb.create_sheet(sh.name[:31])
        for r in range(sh.nrows):
            row: list[object] = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))  # type: ignore[arg-type]  # a date cell holds a float
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(bool(cell.value))
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append(None)
                else:
                    row.append(cell.value)
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _distinct(header: list[str]) -> list[str]:
    """A header with each repeated name numbered, so every column can be named in the register.

    A header that repeats a name, such as Rank and Count once per block, numbers each repeat
    "Count (2)" and on. Blank names stay blank.
    """
    seen: dict[str, int] = {}
    out = []
    for h in header:
        if h:
            seen[h] = seen.get(h, 0) + 1
            out.append(h if seen[h] == 1 else f"{h} ({seen[h]})")
        else:
            out.append(h)
    return out


def read_xlsx(data: bytes, sheet: str, header_row: int) -> pa.Table:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    if sheet and sheet not in wb.sheetnames:
        msg = f"sheet '{sheet}' not in workbook: {wb.sheetnames}"
        raise NormaliseError(msg)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = ws.iter_rows(min_row=header_row, values_only=True)
    header = _distinct([_cell(h).strip() for h in next(rows)])
    while header and not header[-1]:
        header.pop()
    cols: list[list[str]] = [[] for _ in header]
    for row in rows:
        vals = [_cell(v) for v in row[: len(header)]]
        if not any(vals):
            continue
        vals += [""] * (len(header) - len(vals))
        for c, v in zip(cols, vals, strict=True):
            c.append(v)
    wb.close()
    return pa.table({h: pa.array(c, pa.string()) for h, c in zip(header, cols, strict=True)})


def read_wide(data: bytes, ds: Dataset) -> tuple[pa.Table, list[str]]:  # noqa: C901, PLR0912, PLR0915 - one reader for every wide layout, read in order
    """A presentation table as one row per data cell.

    A row is per group of cells instead when the last header row names the fields. Header cells
    left blank beside a filled one are merged cells and take its value, as do row headers named
    in fill_down. A row with no data is a note and is skipped. Values of the last header row that
    no field names are returned as held.
    """
    # normalise reads a wide table only for an entry that declares one.
    w = cast("Wide", ds.wide)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    sheets = w["sheets"] or [ds.source.sheet or wb.sheetnames[0]]
    if w.get("sheet_match"):
        sheets = [n for n in wb.sheetnames if re.search(w["sheet_match"], n)]
        if not sheets:
            msg = f"{ds.slug}: no sheet matches '{w['sheet_match']}' in {wb.sheetnames}"
            raise NormaliseError(msg)
    missing = [n for n in sheets if n not in wb.sheetnames]
    if missing:
        msg = f"{ds.slug}: sheets {missing} not in workbook: {wb.sheetnames}"
        raise NormaliseError(msg)
    measures = {m.group(1): f.source for f in ds.fields if (m := CELL_OF_RE.match(f.source))}
    n, k, first = w["header_rows"], w["row_headers"], w["first_column"] - 1
    levels = n - 1 if measures else n
    out: dict[str, list[str]] = {f.source: [] for f in ds.fields}
    held: list[str] = []
    for name in sheets:
        rows = [
            [_cell(v).strip() for v in r]
            for r in wb[name].iter_rows(min_row=w["header_row"], values_only=True)
        ]
        width = max((len(r) for r in rows), default=0)
        rows = [r + [""] * (width - len(r)) for r in rows]
        head = [r[first + k :] for r in rows[:n]]
        if len(head) < n:
            msg = f"{ds.slug}: sheet '{name}' has fewer than {n} header rows"
            raise NormaliseError(msg)
        for r in head[:-1]:
            for j in range(1, len(r)):
                r[j] = r[j] or r[j - 1]
        groups: dict[tuple[str, ...], dict[str, int]] = {}
        match = re.compile(w["column_match"]) if w.get("column_match") else None
        for j, last in enumerate(head[-1]):
            if not last:
                continue
            if match and not match.search(last):
                if last not in held:
                    held.append(last)
                continue
            if measures:
                if last not in measures:
                    if last not in held:
                        held.append(last)
                    continue
                groups.setdefault(tuple(h[j] for h in head[:levels]), {})[measures[last]] = j
            else:
                groups[tuple(h[j] for h in head)] = {CELL_SOURCE: j}
        if not groups:
            msg = f"{ds.slug}: sheet '{name}' has no columns to read"
            raise NormaliseError(msg)
        found = {src for g in groups.values() for src in g}
        absent = [h for h, src in measures.items() if src not in found]
        if absent:
            msg = f"{ds.slug}: sheet '{name}' has no columns headed {absent}, which the register names"
            raise NormaliseError(msg)
        cols = [j for g in groups.values() for j in g.values()]
        above = [""] * k
        for r in rows[n:]:
            ids, vals = r[first : first + k], r[first + k :]
            if not any(vals[j] for j in cols):
                continue
            for p in w["fill_down"]:
                ids[p - 1] = ids[p - 1] or above[p - 1]
            above = ids
            for key, cells in groups.items():
                for src, col in out.items():
                    if src == SHEET_SOURCE:
                        col.append(name)
                    elif m := ROW_HEADER_RE.match(src):
                        col.append(ids[int(m.group(1)) - 1])
                    elif m := COLUMN_HEADER_RE.match(src):
                        col.append(key[int(m.group(1)) - 1])
                    else:
                        col.append(vals[cells[src]] if src in cells else "")
    wb.close()
    return pa.table({s: pa.array(c, pa.string()) for s, c in out.items()}), held


def read_geojson(data: bytes) -> pa.Table:
    """One row per feature: its properties in first-seen order, then the point's coordinates.

    The longitude and latitude go under LON_SOURCE and LAT_SOURCE. Every cell is text so typing
    follows one path.
    """
    fc = json.loads(data.decode("utf-8-sig"))
    feats = fc.get("features") or []
    header: list[str] = []
    seen: set[str] = set()
    for ft in feats:
        for k in ft.get("properties") or {}:
            if k not in seen:
                seen.add(k)
                header.append(k)
    cols: dict[str, list[str]] = {h: [] for h in [*header, LON_SOURCE, LAT_SOURCE]}
    for ft in feats:
        props = ft.get("properties") or {}
        for h in header:
            cols[h].append(_cell(props.get(h)))
        geom = ft.get("geometry") or {}
        xy = geom.get("coordinates") if geom.get("type") == "Point" else None
        cols[LON_SOURCE].append(_cell(xy[0]) if xy else "")
        cols[LAT_SOURCE].append(_cell(xy[1]) if xy else "")
    return pa.table({h: pa.array(c, pa.string()) for h, c in cols.items()})


XML_JOIN = " | "


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def read_xml(data: bytes, record: str) -> pa.Table:
    """One row per element named `record`, wherever it sits and whatever its namespace.

    Its attributes are columns named @attribute, each child element's text a column named by the
    child, and each child attribute child@attribute. A child that repeats within a record gives
    its values in document order joined by XML_JOIN. Every cell is text.
    """
    root = ET.fromstring(data.decode("utf-8-sig").encode("utf-8"))  # noqa: S314 - Expat refuses entity expansion and loads no external entity
    records = [e for e in root.iter() if _local(e.tag) == record]
    if not records:
        msg = f"the XML holds no <{record}> element"
        raise NormaliseError(msg)
    # A record inside a record would be read twice, once as a row and once as a child.
    if any(_local(x.tag) == record for e in records for x in e.iter() if x is not e):
        msg = f"a <{record}> element holds another <{record}>; name the outer one"
        raise NormaliseError(msg)
    header: list[str] = []
    rows: list[dict[str, list[str]]] = []
    for e in records:
        row: dict[str, list[str]] = {}
        for k, v in e.attrib.items():
            row.setdefault(f"@{_local(k)}", []).append(v)
        for child in e:
            name = _local(child.tag)
            row.setdefault(name, []).append((child.text or "").strip())
            for k, v in child.attrib.items():
                row.setdefault(f"{name}@{_local(k)}", []).append(v)
        for k in row:
            if k not in header:
                header.append(k)
        rows.append(row)
    return pa.table(
        {h: pa.array([XML_JOIN.join(r.get(h, [])) for r in rows], pa.string()) for h in header}
    )


def unwrap(data: bytes, ext: str, member: str) -> tuple[bytes, str]:
    """The data file out of a zip the publisher serves: the named member, or the only one."""
    if ext != "zip":
        return data, ext
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = [n for n in z.namelist() if not n.endswith("/")]
        if member:
            if member not in names:
                msg = f"'{member}' is not in the zip: {names}"
                raise NormaliseError(msg)
            name = member
        else:
            data_names = [
                n
                for n in names
                if n.lower().endswith(
                    (".csv", ".tsv", ".xlsx", ".xls", ".geojson", ".json", ".xml")
                )
            ]
            if len(data_names) != 1:
                msg = f"the zip holds {len(data_names)} data files; name one: {names}"
                raise NormaliseError(msg)
            name = data_names[0]
        return z.read(name), name.rsplit(".", 1)[-1].lower()


DATE_HEADER = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def unpivot(raw: pa.Table, ds: Dataset) -> tuple[pa.Table, list[str]]:
    """One row per identifying row and dated column, ordered by the source row then the column.

    Headers that are not dates are returned as held, like any column the register does not name.
    """
    keep = [f.source for f in ds.fields if f.source not in (HEADER_SOURCE, CELL_SOURCE)]
    names = {c.strip(): c for c in raw.column_names}
    missing = [k for k in keep if k not in names]
    if missing:
        msg = f"{ds.slug}: columns named in the register are absent upstream: {missing}"
        raise NormaliseError(msg)
    wide = [c for c in raw.column_names if c.strip() not in keep]
    dated = [c for c in wide if DATE_HEADER.match(c.strip())]
    held = [c for c in wide if not DATE_HEADER.match(c.strip())]
    if not dated:
        msg = f"{ds.slug}: no dated columns to unpivot; the header format may have changed"
        raise NormaliseError(msg)
    n, m = raw.num_rows, len(dated)
    order = pa.array([j * n + i for i in range(n) for j in range(m)], pa.int64())
    cols = {
        k: pa.concat_arrays([raw.column(names[k]).combine_chunks()] * m).take(order) for k in keep
    }
    cols[HEADER_SOURCE] = pa.concat_arrays(
        [pa.array([c.strip()] * n, pa.string()) for c in dated]
    ).take(order)
    cols[CELL_SOURCE] = pa.concat_arrays([raw.column(c).combine_chunks() for c in dated]).take(
        order
    )
    return pa.table(cols), held


def _blank_to_null(arr: Arr) -> Arr:
    arr = pc.utf8_trim_whitespace(arr)
    return pc.if_else(pc.equal(arr, ""), pa.scalar(None, pa.string()), arr)  # type: ignore[call-overload, no-any-return]  # pyarrow-stubs 20 takes no Python scalar


def _examples(arr: Arr, mask: Arr, n: int = 5) -> list[str]:
    return [str(v) for v in pc.filter(arr, mask).slice(0, n).to_pylist()]


THOUSANDS = r"^-?\d{1,3}(,\d{3})+(\.\d+)?$"


def _plain_number(arr: Arr) -> Arr:
    """A figure a publisher wrote with thousands separators, 19,918, as 19918.

    Only a cell that is wholly such a figure is touched.
    """
    grouped = pc.fill_null(pc.match_substring_regex(arr, THOUSANDS), fill_value=False)  # type: ignore[call-arg]  # pyarrow-stubs 20 types fill_null as coalesce
    if not pc.any(grouped).as_py():
        return arr
    return pc.if_else(grouped, pc.replace_substring(arr, ",", ""), arr)  # type: ignore[no-any-return]  # pyarrow-stubs 20 leaves if_else untyped


def _strptime(arr: Arr, f: Field) -> Arr:
    """Dates in the field's format, or in any of several formats joined by | in one string.

    A publisher's workbooks may hold a date as a date in one and as text in another. A value in
    none of the formats stops the build.
    """
    formats = f.date_format.split("|")
    if len(formats) == 1:
        return pc.strptime(arr, format=formats[0], unit="s")
    out: Arr | None = None
    for fmt in formats:
        got = pc.strptime(arr, format=fmt, unit="s", error_is_null=True)
        out = got if out is None else pc.coalesce(out, got)
    bad = pc.and_(pc.is_valid(arr), pc.is_null(out))  # type: ignore[arg-type]  # formats is never empty, so out is set
    if pc.any(bad).as_py():
        msg = f"{f.name}: values in none of the formats {formats}: {_examples(arr, bad)}"
        raise NormaliseError(msg)
    return out  # type: ignore[return-value]  # formats is never empty, so out is set


def convert(arr: Arr, f: Field, suppression: tuple[str, ...]) -> tuple[Arr, Arr | None]:  # noqa: C901 - one branch per field type
    """Return (typed array, suppression mask or None)."""
    arr = _blank_to_null(arr)
    if f.null_values:
        unknown = pc.fill_null(
            pc.is_in(arr, value_set=pa.array(list(f.null_values))),
            fill_value=False,  # type: ignore[call-arg]  # pyarrow-stubs 20 types fill_null as coalesce
        )
        arr = pc.if_else(unknown, pa.scalar(None, pa.string()), arr)
    sup = None
    if suppression and f.type in ("integer", "number"):
        sup = pc.fill_null(pc.is_in(arr, value_set=pa.array(list(suppression))), fill_value=False)  # type: ignore[call-arg]  # pyarrow-stubs 20 types fill_null as coalesce
        if pc.any(sup).as_py():
            arr = pc.if_else(sup, pa.scalar(None, pa.string()), arr)
        else:
            sup = None
    try:
        if f.type == "string":
            return arr, sup
        if f.type in ("integer", "number"):
            return pc.cast(_plain_number(arr), ARROW_TYPES[f.type]), sup
        if f.type == "boolean":
            t = pc.fill_null(
                pc.is_in(arr, value_set=pa.array(list(f.true_values))),
                fill_value=False,  # type: ignore[call-arg]  # pyarrow-stubs 20 types fill_null as coalesce
            )
            fl = pc.fill_null(
                pc.is_in(arr, value_set=pa.array(list(f.false_values))),
                fill_value=False,  # type: ignore[call-arg]  # pyarrow-stubs 20 types fill_null as coalesce
            )
            bad = pc.and_(pc.is_valid(arr), pc.invert(pc.or_(t, fl)))
            if pc.any(bad).as_py():
                msg = f"{f.name}: values outside true/false sets: {_examples(arr, bad)}"
                raise NormaliseError(msg)
            return pc.if_else(t, True, pc.if_else(fl, False, pa.scalar(None, pa.bool_()))), sup  # noqa: FBT003 - pyarrow's if_else takes them by position
        if f.type in ("date", "datetime"):
            if f.date_format == "epoch_ms":
                # An ArcGIS service gives dates as milliseconds since 1970, UTC.
                ms = pc.cast(arr, pa.int64())
                ts = pc.cast(pc.cast(ms, pa.timestamp("ms")), pa.timestamp("s"))
            else:
                ts = _strptime(arr, f)
            return (pc.cast(ts, pa.date32()) if f.type == "date" else ts), sup
    except pa.ArrowInvalid as e:
        msg = f"{f.name}: cannot type as {f.type}: {e}"
        raise NormaliseError(msg) from e
    msg = f"{f.name}: unknown type {f.type}"
    raise NormaliseError(msg)


def _is_shapefile(ext: str, member: str) -> bool:
    """A layer GDAL reads, as opposed to a table: a shapefile or GeoPackage, or one inside a zip."""
    return ext in ("shp", "gpkg") or (ext == "zip" and member.lower().endswith((".shp", ".gpkg")))


def normalise(ds: Dataset, m: Manifest, data: bytes) -> Table:  # noqa: C901, PLR0912, PLR0915 - the normalising steps, read in order
    held: list[str] = []
    short: list[int] = []
    shapes = None
    if ds.geometry and ds.geometry["kind"] != "point":
        raw, shapes = read_shapes(data, m.ext, ds.source.member, ds.geometry["crs"])
    elif ds.geometry and _is_shapefile(m.ext, ds.source.member):
        raw = read_points(data, m.ext, ds.source.member)
    else:
        data, ext = unwrap(data, m.ext, ds.source.member)
        if ext == "xls":
            data, ext = xls_to_xlsx(data), "xlsx"
        if ds.wide:
            if ext not in ("xlsx", "xlsm"):
                msg = f"{ds.slug}: a wide table is read from a workbook, got {ext}"
                raise NormaliseError(msg)
            raw, held = read_wide(data, ds)
        elif ext in ("xlsx", "xlsm"):
            raw = read_xlsx(data, ds.source.sheet, ds.source.header_row)
        elif ext in ("geojson", "json"):
            raw = read_geojson(data)
        elif ext == "xml":
            if not ds.source.record:
                msg = f"{ds.slug}: an XML source names its record element in source.record"
                raise NormaliseError(msg)
            raw = read_xml(data, ds.source.record)
        else:
            enc = m.encoding if m.ext != "zip" else detect_encoding(data, ds.source.encoding)
            raw = read_csv(data, enc, ds.source.delimiter, ds.source.header_row, short)
    if ds.unpivot:
        raw, held = unpivot(raw, ds)

    upstream = {c.strip(): c for c in raw.column_names}
    own = [f for f in ds.fields if not is_spine(f.source)]
    missing = [f.source for f in own if f.source not in upstream]
    if missing:
        msg = f"{ds.slug}: columns named in the register are absent upstream: {missing}"
        raise NormaliseError(msg)
    declared = {f.source for f in own}
    extra = [c for c in upstream if c not in declared] + held
    unknown = [c for c in extra if c not in ds.omit]
    omitted = [c for c in extra if c in ds.omit]
    cols: list[Arr] = []
    names: list[str] = []
    masks: dict[str, Arr] = {}
    for f in own:
        arr, sup = convert(raw.column(upstream[f.source]), f, ds.suppression)
        cols.append(arr)
        names.append(f.name)
        if sup is not None:
            masks[f.name] = sup
    suppressed = 0
    if masks:
        flags: list[list[str]] = [[] for _ in range(raw.num_rows)]
        for name, mask in masks.items():
            for i, v in enumerate(mask.to_pylist()):
                if v:
                    flags[i].append(name)
                    suppressed += 1
        cols.append(pa.array(flags, pa.list_(pa.string())))
        names.append("suppressed")
    table = pa.table(cols, names=names)
    return Table(
        dataset=ds,
        manifest=m,
        table=table,
        unknown_columns=unknown,
        suppressed_cells=suppressed,
        short_rows=sum(short),
        omitted_columns=omitted,
        geometry=shapes,
    )
